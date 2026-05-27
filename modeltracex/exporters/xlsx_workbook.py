"""XLSX lineage workbook (Spec Part B) — exactly 9 sheets, fixed order + columns.

``openpyxl``. The sheet names and column headers in ``SHEETS`` are the Part B
contract (the golden test asserts against them). Provenance (E/H/I/U), Confidence,
and Severity cells are conditionally color-coded; ``Table ID`` is a consistent FK
across sheets for spreadsheet-side joins.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from modeltracex.state import RunState

# (sheet name, [exact column headers]) — Part B order is the contract.
SHEETS: list[tuple[str, list[str]]] = [
    (
        "Run Summary",
        [
            "Run ID",
            "Timestamp",
            "Tool version",
            "LLM provider / model",
            "Languages detected",
            "# Models",
            "# Input tables",
            "# Output tables",
            "# Table-level edges",
            "# Column-level edges",
            "# DQ rules",
            "Tokens",
            "Est. cost",
        ],
    ),
    (
        "Model Inventory",
        [
            "Model ID",
            "Model label",
            "Language",
            "Source file(s)",
            "Purpose",
            "# Inputs",
            "# Outputs",
            "Confidence",
            "Status",
        ],
    ),
    (
        "Tables",
        [
            "Table ID",
            "Table name",
            "Role",
            "Source system",
            "Produced by model(s)",
            "Consumed by model(s)",
            "Grain",
            "Source",
        ],
    ),
    ("Columns", ["Table ID", "Column name", "Inferred type", "Role", "Used in", "Source"]),
    (
        "Table Lineage",
        [
            "Edge ID",
            "Source table ID",
            "Target table ID",
            "Model ID",
            "Transformation type",
            "Notes",
            "Source",
            "Confidence",
        ],
    ),
    (
        "Column Lineage",
        [
            "Edge ID",
            "Source table.column",
            "Target table.column",
            "Model ID",
            "Transformation type",
            "Transformation expression",
            "Join key(s)",
            "Source",
            "Confidence",
        ],
    ),
    (
        "Data Quality Rules",
        [
            "Rule ID",
            "Table.column",
            "DQ dimension",
            "Rule",
            "Rationale / usage",
            "Code evidence",
            "Suggested severity",
            "Source",
            "Status",
        ],
    ),
    ("Source Systems", ["Source system", "# Tables", "Table IDs"]),
    ("Issues & Confidence Log", ["Model ID", "Severity", "Message"]),
]

_FILL = {
    "E": "DDEBF7",
    "H": "E2EFDA",
    "I": "FCE4D6",
    "U": "FFF2CC",
    "High": "C6EFCE",
    "Medium": "FFEB9C",
    "Low": "FFC7CE",
}
_COLORED_HEADERS = {"Source", "Confidence", "Suggested severity"}


class XlsxExporter:
    fmt = "xlsx"

    def export(self, state: RunState, out_dir: str) -> list[str]:
        rows = self._rows(state)
        wb = Workbook()
        wb.remove(wb.active)
        for name, headers in SHEETS:
            ws = wb.create_sheet(title=name)
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            colored_cols = [i for i, h in enumerate(headers) if h in _COLORED_HEADERS]
            for row in rows[name]:
                ws.append(row)
                for idx in colored_cols:
                    fill = _FILL.get(str(row[idx]))
                    if fill:
                        ws.cell(row=ws.max_row, column=idx + 1).fill = PatternFill(
                            "solid", fgColor=fill
                        )
        path = Path(out_dir) / "lineage_workbook.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        return [str(path)]

    def _rows(self, state: RunState) -> dict[str, list[list[object]]]:
        inputs = sum(1 for t in state.tables if t.consumed_by)
        outputs = sum(1 for t in state.tables if t.produced_by)
        langs = ", ".join(sorted({m.language.value for m in state.models}))
        return {
            "Run Summary": [
                [
                    state.run.run_id,
                    state.run.timestamp,
                    state.run.tool_version,
                    f"{state.run.llm_provider} / {state.run.llm_model}",
                    langs,
                    len(state.models),
                    inputs,
                    outputs,
                    len(state.table_edges),
                    len(state.column_edges),
                    len(state.dq_rules),
                    state.run.tokens,
                    state.run.est_cost,
                ]
            ],
            "Model Inventory": [
                [
                    m.model_id,
                    m.label,
                    m.language.value,
                    ", ".join(m.source_files),
                    m.purpose,
                    sum(1 for t in state.tables if m.model_id in t.consumed_by),
                    sum(1 for t in state.tables if m.model_id in t.produced_by),
                    m.confidence.value,
                    m.status.value,
                ]
                for m in state.models
            ],
            "Tables": [
                [
                    t.table_id,
                    t.name,
                    t.role.value,
                    t.source_system or "",
                    ", ".join(t.produced_by),
                    ", ".join(t.consumed_by),
                    t.grain or "",
                    t.source.value,
                ]
                for t in state.tables
            ],
            "Columns": [
                [
                    t.table_id,
                    c.name,
                    c.inferred_type or "",
                    c.role.value,
                    ", ".join(c.used_in),
                    c.source.value,
                ]
                for t in state.tables
                for c in t.columns
            ],
            "Table Lineage": [
                [
                    e.edge_id,
                    e.source_table_id,
                    e.target_table_id,
                    e.model_id,
                    e.transformation_type.value,
                    e.notes,
                    e.source.value,
                    e.confidence.value,
                ]
                for e in state.table_edges
            ],
            "Column Lineage": [
                [
                    e.edge_id,
                    e.source_element,
                    e.target_element,
                    e.model_id,
                    e.transformation_type.value,
                    e.expression or "",
                    ", ".join(e.join_keys),
                    e.source.value,
                    e.confidence.value,
                ]
                for e in state.column_edges
            ],
            "Data Quality Rules": [
                [
                    r.rule_id,
                    r.element,
                    r.dimension.value,
                    r.rule_statement,
                    r.rationale,
                    r.code_evidence,
                    r.severity.value,
                    r.source.value,
                    r.status.value,
                ]
                for r in state.dq_rules
            ],
            "Source Systems": [
                [s.name, len(s.tables), ", ".join(s.tables)] for s in state.source_systems
            ],
            "Issues & Confidence Log": [
                [i.model_id or "", i.severity, i.message] for i in state.issues
            ],
        }


__all__ = ["XlsxExporter", "SHEETS"]
