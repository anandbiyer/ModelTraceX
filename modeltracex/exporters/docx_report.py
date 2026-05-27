"""DOCX model documentation (Spec Part A) — one doc per model + a project summary.

Each of the 13 sections is a pure projection of ``RunState`` (FR-4.2). Empty
sections render the explicit "Not identified from code" string (Part A rule)
rather than being omitted, so the structure is repeatable. Inputs, outputs, and
DQ rules render as tables, not prose.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument

from modeltracex.exporters.base import NOT_IDENTIFIED, safe_filename
from modeltracex.lineage.graph import LineageGraph
from modeltracex.state import DQRule, ModelDoc, RunState, Table, TableEdge

# The 13 Part A sections, in order — also the golden-test contract.
SECTION_TITLES = [
    "Cover / metadata",
    "Executive summary",
    "Purpose and scope",
    "Data inputs",
    "Assumptions and dependencies",
    "Methodology / processing logic",
    "Calculation logic",
    "Data outputs",
    "Data lineage summary",
    "Data quality considerations",
    "Limitations and risks",
    "Provenance and confidence",
    "Appendix: source listing",
]


def _table(doc: DocxDocument, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = str(value)


def _para_or_missing(doc: DocxDocument, text: str) -> None:
    doc.add_paragraph(text if text.strip() else NOT_IDENTIFIED)


def _bullets_or_missing(doc: DocxDocument, items: list[str]) -> None:
    if not items:
        doc.add_paragraph(NOT_IDENTIFIED)
        return
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


class DocxExporter:
    fmt = "docx"

    def export(self, state: RunState, out_dir: str) -> list[str]:
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        usage_by_model: dict[str, set[str]] = {m.model_id: set() for m in state.models}
        for u in state.usage_observations:
            usage_by_model.setdefault(u.model_id, set()).add(u.element)

        paths: list[str] = []
        for model in state.models:
            paths.append(self._model_doc(state, model, usage_by_model, out_dir))
        paths.append(self._project_summary(state, out_dir))
        return paths

    def _model_doc(
        self, state: RunState, model: ModelDoc, usage_by_model: dict[str, set[str]], out_dir: str
    ) -> str:
        doc = Document()
        doc.add_heading(f"Model documentation — {model.label}", level=0)
        inputs = [t for t in state.tables if model.model_id in t.consumed_by]
        outputs = [t for t in state.tables if model.model_id in t.produced_by]
        edges = [e for e in state.table_edges if e.model_id == model.model_id]
        elements = usage_by_model.get(model.model_id, set())
        rules = [r for r in state.dq_rules if r.element in elements]

        for title in SECTION_TITLES:
            doc.add_heading(title, level=1)
            self._section_body(doc, title, state, model, inputs, outputs, edges, rules)
        path = Path(out_dir) / f"{safe_filename(model.label)}.docx"
        doc.save(str(path))
        return str(path)

    def _section_body(  # noqa: PLR0913
        self,
        doc: DocxDocument,
        title: str,
        state: RunState,
        model: ModelDoc,
        inputs: list[Table],
        outputs: list[Table],
        edges: list[TableEdge],
        rules: list[DQRule],
    ) -> None:
        if title == "Cover / metadata":
            _table(
                doc,
                ["Field", "Value"],
                [
                    ["Model label", model.label],
                    ["Language", model.language.value],
                    ["Source file(s)", ", ".join(model.source_files) or NOT_IDENTIFIED],
                    ["Status", model.status.value],
                    ["Analysis timestamp", state.run.timestamp],
                    ["Run ID", state.run.run_id],
                    ["Tool version", state.run.tool_version],
                    ["Provider / model", f"{state.run.llm_provider} / {state.run.llm_model}"],
                ],
            )
        elif title == "Executive summary":
            _para_or_missing(doc, model.executive_summary)
        elif title == "Purpose and scope":
            _para_or_missing(doc, model.purpose)
        elif title == "Data inputs":
            if inputs:
                _table(
                    doc,
                    ["Table", "Source system", "Grain", "Columns"],
                    [
                        [
                            t.name,
                            t.source_system or "",
                            t.grain or "",
                            ", ".join(c.name for c in t.columns),
                        ]
                        for t in inputs
                    ],
                )
            else:
                doc.add_paragraph(NOT_IDENTIFIED)
        elif title == "Assumptions and dependencies":
            _bullets_or_missing(doc, model.assumptions)
        elif title == "Methodology / processing logic":
            if model.methodology_steps:
                for i, step in enumerate(model.methodology_steps, 1):
                    doc.add_paragraph(f"{i}. {step}")
            else:
                doc.add_paragraph(NOT_IDENTIFIED)
        elif title == "Calculation logic":
            if model.calculations:
                _table(
                    doc,
                    ["Target", "Expression"],
                    [[c.target, c.expression] for c in model.calculations],
                )
            else:
                doc.add_paragraph(NOT_IDENTIFIED)
        elif title == "Data outputs":
            if outputs:
                _table(
                    doc,
                    ["Table", "Grain", "Columns"],
                    [
                        [t.name, t.grain or "", ", ".join(c.name for c in t.columns)]
                        for t in outputs
                    ],
                )
            else:
                doc.add_paragraph(NOT_IDENTIFIED)
        elif title == "Data lineage summary":
            if edges:
                names = {t.table_id: t.name for t in state.tables}
                doc.add_paragraph(
                    "End-to-end flow for this model (see the lineage workbook for full detail):"
                )
                for e in edges:
                    doc.add_paragraph(
                        f"{names.get(e.source_table_id, e.source_table_id)} -> "
                        f"{names.get(e.target_table_id, e.target_table_id)} "
                        f"({e.transformation_type.value})",
                        style="List Bullet",
                    )
            else:
                doc.add_paragraph(NOT_IDENTIFIED)
        elif title == "Data quality considerations":
            if rules:
                _table(
                    doc,
                    ["DQ dimension", "Element", "Severity", "Rule"],
                    [
                        [r.dimension.value, r.element, r.severity.value, r.rule_statement]
                        for r in sorted(rules, key=lambda r: r.dimension.value)
                    ],
                )
            else:
                doc.add_paragraph(NOT_IDENTIFIED)
        elif title == "Limitations and risks":
            _bullets_or_missing(doc, model.limitations)
        elif title == "Provenance and confidence":
            doc.add_paragraph(f"Overall model confidence: {model.confidence.value}.")
            doc.add_paragraph(
                "Facts are tagged by origin: E = LLM-extracted, H = heuristic/parser, "
                "I = LLM-inferred, U = user override. See the workbook Source columns."
            )
            applied = list(state.overrides)
            if applied:
                doc.add_paragraph("User overrides applied:")
                for o in applied:
                    doc.add_paragraph(f"{o.target}.{o.field} -> {o.new}", style="List Bullet")
        elif title == "Appendix: source listing":
            note = "Source listing honors the run's retention mode. Files analyzed: " + (
                ", ".join(model.source_files) or NOT_IDENTIFIED
            )
            doc.add_paragraph(note)

    def _project_summary(self, state: RunState, out_dir: str) -> str:
        doc = Document()
        doc.add_heading("Project summary", level=0)

        doc.add_heading("Model inventory", level=1)
        _table(
            doc,
            ["Model", "Language", "Status", "Confidence", "Purpose"],
            [
                [m.label, m.language.value, m.status.value, m.confidence.value, m.purpose]
                for m in state.models
            ],
        )

        doc.add_heading("Cross-model lineage", level=1)
        graph = LineageGraph(state)
        paths = graph.source_to_output_paths()
        if paths:
            for node_path in paths:
                doc.add_paragraph(
                    " -> ".join(graph.name_of(n) for n in node_path), style="List Bullet"
                )
        else:
            doc.add_paragraph(NOT_IDENTIFIED)

        doc.add_heading("Data quality register", level=1)
        if state.dq_rules:
            _table(
                doc,
                ["Rule ID", "Element", "Dimension", "Severity", "Status"],
                [
                    [r.rule_id, r.element, r.dimension.value, r.severity.value, r.status.value]
                    for r in state.dq_rules
                ],
            )
        else:
            doc.add_paragraph(NOT_IDENTIFIED)

        doc.add_heading("Source systems and outputs", level=1)
        doc.add_paragraph(
            "Source systems: " + (", ".join(s.name for s in state.source_systems) or NOT_IDENTIFIED)
        )
        outputs = [t.name for t in state.tables if t.role.value == "Output"]
        doc.add_paragraph("Outputs: " + (", ".join(outputs) or NOT_IDENTIFIED))

        path = Path(out_dir) / "project_summary.docx"
        doc.save(str(path))
        return str(path)


__all__ = ["DocxExporter", "SECTION_TITLES"]
