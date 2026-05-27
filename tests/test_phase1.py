"""Phase 1 — Headless E2E tests (P1-T1 .. P1-T8).

Offline (FakeProvider / adapter scans only). Covers adapters, the DQ Part-C
mapper, chunk-merge idempotence, cross-model stitch, partial-failure isolation,
the DOCX/XLSX/CSV exporter goldens, ingestion formats, and the Phase-1 exit E2E
(a SAS+Python project → valid Part A DOCX + Part B XLSX).
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

import modeltracex.adapters  # noqa: F401  (register built-ins)
from modeltracex.adapters.base import ADAPTERS
from modeltracex.analysis.chunking import chunk_code, estimate_tokens
from modeltracex.analysis.orchestrator import ModelInput, analyze_run
from modeltracex.dq.engine import infer_rules
from modeltracex.dq.rules_table import PART_C
from modeltracex.exporters.base import NOT_IDENTIFIED
from modeltracex.exporters.csv_compat import HEADERS as CSV_HEADERS
from modeltracex.exporters.csv_compat import CsvCompatExporter
from modeltracex.exporters.docx_report import SECTION_TITLES, DocxExporter
from modeltracex.exporters.xlsx_workbook import SHEETS, XlsxExporter
from modeltracex.llm.fake import FakeProvider
from modeltracex.state import (
    Language,
    ModelStatus,
    Provenance,
    TableRole,
    UsageKind,
    UsageObservation,
)

FIXTURES = Path(__file__).parent / "fixtures"
CORPUS = FIXTURES / "corpus"
DQ = CORPUS / "dq_coverage"
INGESTION = CORPUS / "ingestion"
SAS = ADAPTERS[Language.SAS]
PY = ADAPTERS[Language.PYTHON]


# --------------------------------------------------------------------------- #
# P1-T1 — Adapter StructuralScan (per corpus file) + SAS v1 golden-master parity
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    "kind_dir", sorted(d for d in DQ.glob("*") if d.is_dir()), ids=lambda p: p.name
)
def test_adapter_emits_expected_usage(kind_dir: Path) -> None:
    expected = json.loads((kind_dir / "expected.json").read_text("utf-8"))
    want_kind = expected["usage_kind"]
    want_col = expected["expected_usages"][0]["element"].split(".")[-1]
    for adapter, fname in [(SAS, "model.sas"), (PY, "model.py")]:
        scan = adapter.scan((kind_dir / fname).read_text("utf-8"), "m")
        emitted = {(u.element, u.usage_kind.value) for u in scan.usages}
        assert (want_col, want_kind) in emitted, (
            f"{fname}: missing ({want_col},{want_kind}) in {emitted}"
        )


def test_sas_v1_golden_master_parity() -> None:
    facts = json.loads((CORPUS / "golden_v1" / "expected.json").read_text("utf-8"))["v1_facts"]
    scan = SAS.scan((CORPUS / "golden_v1" / "v1_sample.sas").read_text("utf-8"), "m")
    assert sorted(scan.inputs) == sorted(facts["inputs"])
    assert sorted(scan.outputs) == sorted(facts["outputs"])
    # v2 improvement: emits DQ usages v1 never did.
    assert {u.usage_kind.value for u in scan.usages} >= {
        "denominator",
        "range_filter",
        "aggregated",
    }


def test_python_adapter_inputs_outputs_and_split_points() -> None:
    exp = json.loads((CORPUS / "languages" / "python" / "expected.json").read_text("utf-8"))[
        "fixtures"
    ]
    for fname, e in exp.items():
        scan = PY.scan((CORPUS / "languages" / "python" / fname).read_text("utf-8"), "m")
        assert sorted(scan.inputs) == sorted(e["inputs"])
        assert sorted(scan.outputs) == sorted(e["outputs"])
        if "split_point_defs" in e:
            assert len(scan.split_points) == len(e["split_point_defs"])


# --------------------------------------------------------------------------- #
# P1-T2 — DQ engine Part-C mapping + coverage gate
# --------------------------------------------------------------------------- #
def test_rules_table_matches_corpus_part_c() -> None:
    fixture = json.loads((FIXTURES / "part_c_dq_table.json").read_text("utf-8"))
    for row in fixture["rules"]:
        rule = PART_C[UsageKind(row["usage_kind"])]
        assert rule.dimension.value == row["dimension"]
        assert rule.severity.value == row["severity"]
        assert rule.statement == row["statement"]
        if "also_dimension" in row:
            assert rule.also_dimension is not None
            assert rule.also_dimension.value == row["also_dimension"]


def test_dq_engine_maps_every_usage_kind_and_covers_all_dimensions() -> None:
    usages = [
        UsageObservation(
            element=f"t.{k.value}",
            usage_kind=k,
            evidence="ev",
            model_id="m",
            source=Provenance.HEURISTIC,
        )
        for k in PART_C
    ]
    rules = infer_rules(usages)
    by_element = {(r.element, r.dimension.value) for r in rules}
    # join_key yields two rules (Uniqueness + Consistency).
    assert ("t.join_key", "Uniqueness") in by_element
    assert ("t.join_key", "Consistency") in by_element
    assert {r.dimension.value for r in rules} == {
        "Completeness",
        "Validity",
        "Uniqueness",
        "Consistency",
        "Accuracy",
        "Timeliness",
    }
    assert all(r.source is Provenance.HEURISTIC for r in rules)


# --------------------------------------------------------------------------- #
# P1-T3 — Chunk-merge idempotence
# --------------------------------------------------------------------------- #
def test_chunking_preserves_content() -> None:
    code = (CORPUS / "chunking" / "oversized_etl.sas").read_text("utf-8")
    pts = SAS.split_points(code)
    chunks = chunk_code(code, pts, max_tokens=50)
    assert len(chunks) > 1
    assert "".join(chunks) == code


def _projection(state) -> tuple:
    return (
        sorted((t.table_id, t.role.value) for t in state.tables),
        sorted(e.edge_id for e in state.table_edges),
        sorted(e.edge_id for e in state.column_edges),
        sorted(r.rule_id for r in state.dq_rules),
    )


def test_chunked_and_unchunked_runs_are_identical() -> None:
    code = (CORPUS / "chunking" / "oversized_etl.sas").read_text("utf-8")
    fixed = json.dumps(
        {
            "purpose": "ETL chain",
            "input_tables": [{"name": "raw.source", "columns": ["amount"]}],
            "output_tables": [{"name": "report.final", "columns": ["rank"]}],
            "lineage_rows": ["raw.source.amount -> report.final.rank"],
        }
    )
    provider = FakeProvider(responses=fixed)
    mi = ModelInput(label="etl", language=Language.SAS, source_files=["etl.sas"], code=code)

    whole = analyze_run(provider, [mi], run_id="r", max_chunk_tokens=100_000)
    chunked = analyze_run(provider, [mi], run_id="r", max_chunk_tokens=20)
    assert estimate_tokens(code) > 20  # the second run really did chunk
    assert _projection(whole) == _projection(chunked)


# --------------------------------------------------------------------------- #
# P1-T4 — Cross-model stitch
# --------------------------------------------------------------------------- #
def _stitch_state():
    m1 = json.dumps(
        {
            "purpose": "Build staging.",
            "input_tables": [{"name": "raw.customer_events", "columns": ["cust_id"]}],
            "output_tables": [{"name": "work.staging", "columns": ["cust_id", "amount_net"]}],
            "lineage_rows": ["raw.customer_events.cust_id -> work.staging.cust_id"],
        }
    )
    m2 = json.dumps(
        {
            "purpose": "Score.",
            "input_tables": [{"name": "work.staging", "columns": ["amount_net"]}],
            "output_tables": [{"name": "mart.customer_scores", "columns": ["total_net"]}],
            "lineage_rows": ["work.staging.amount_net -> mart.customer_scores.total_net"],
        }
    )
    provider = FakeProvider(responses={"m1_build_staging": m1, "m2_score": m2})
    inputs = [
        ModelInput(label="m1_build_staging", language=Language.SAS, source_files=["m1.sas"]),
        ModelInput(label="m2_score", language=Language.SAS, source_files=["m2.sas"]),
    ]
    return analyze_run(provider, inputs, run_id="stitch")


def test_cross_model_stitch_makes_one_intermediate_node() -> None:
    from modeltracex.lineage.graph import LineageGraph

    state = _stitch_state()
    staging = next(t for t in state.tables if t.name == "work.staging")
    assert staging.role is TableRole.INTERMEDIATE
    assert staging.produced_by and staging.consumed_by  # written by m1, read by m2
    paths = LineageGraph(state).source_to_output_paths()
    names = [[LineageGraph(state).name_of(n) for n in p] for p in paths]
    assert ["raw.customer_events", "work.staging", "mart.customer_scores"] in names


# --------------------------------------------------------------------------- #
# P1-T5 — Partial-failure isolation
# --------------------------------------------------------------------------- #
def test_partial_failure_isolates_one_model() -> None:
    good = json.dumps({"purpose": "ok", "output_tables": [{"name": "work.clean"}]})
    provider = FakeProvider(responses={"m_good": good, "m_bad": "42"})
    inputs = [
        ModelInput(label="m_good", language=Language.SAS, source_files=["g.sas"]),
        ModelInput(label="m_bad", language=Language.SAS, source_files=["b.sas"]),
    ]
    state = analyze_run(provider, inputs, run_id="pf")
    status = {m.label: m.status for m in state.models}
    assert status["m_bad"] is ModelStatus.FAILED
    assert status["m_good"] is ModelStatus.ANALYZED
    assert any(
        i.model_id and "m_bad" in (i.model_id or "") or i.severity == "error" for i in state.issues
    )
    assert len(state.models) == 2  # batch completed


# --------------------------------------------------------------------------- #
# P1-T6 — DOCX golden (13 sections, "Not identified", tables)
# --------------------------------------------------------------------------- #
def test_docx_has_13_sections_with_missing_rule(tmp_path: Path) -> None:
    state = _stitch_state()
    paths = DocxExporter().export(state, str(tmp_path))
    model_docs = [p for p in paths if not p.endswith("project_summary.docx")]
    doc = Document(model_docs[0])
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == SECTION_TITLES  # all 13, in order
    assert any(NOT_IDENTIFIED in p.text for p in doc.paragraphs)  # empty sections explicit
    assert doc.tables  # inputs/outputs rendered as tables
    assert Path(model_docs[0]).exists()


# --------------------------------------------------------------------------- #
# P1-T7 — XLSX golden (9 sheets, exact headers, color) + legacy CSV
# --------------------------------------------------------------------------- #
def test_xlsx_has_exactly_nine_sheets_with_exact_headers(tmp_path: Path) -> None:
    state = _stitch_state()
    path = XlsxExporter().export(state, str(tmp_path))[0]
    wb = load_workbook(path)
    assert wb.sheetnames == [name for name, _ in SHEETS]
    for name, headers in SHEETS:
        ws = wb[name]
        assert [c.value for c in ws[1]] == headers
    # provenance color-coding present on the Tables sheet's Source column.
    tables_ws = wb["Tables"]
    source_col = [c.value for c in tables_ws[1]].index("Source") + 1
    assert tables_ws.cell(row=2, column=source_col).fill.fgColor.rgb is not None


def test_csv_compat_is_five_columns(tmp_path: Path) -> None:
    state = _stitch_state()
    path = CsvCompatExporter().export(state, str(tmp_path))[0]
    first = Path(path).read_text("utf-8").splitlines()[0].split(",")
    assert first == CSV_HEADERS
    assert len(CSV_HEADERS) == 5


# --------------------------------------------------------------------------- #
# P1-T8 — Ingestion formats + Phase-1 exit E2E
# --------------------------------------------------------------------------- #
@pytest.fixture(scope="module")
def _binaries() -> dict[str, str]:
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "_bc", Path(__file__).parent / "tools" / "build_corpus.py"
    )
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.materialize()


def test_ingestion_formats_yield_equivalent_models(_binaries: dict[str, str]) -> None:
    from modeltracex.ingestion import assemble_models, read_path, unzip_project

    def scan_of(models) -> tuple:
        m = models[0]
        scan = SAS.scan(m.code, "m")
        return (tuple(sorted(scan.inputs)), tuple(sorted(scan.outputs)))

    single = assemble_models([read_path(INGESTION / "single" / "model.sas")], one_model=True)
    paste = assemble_models([read_path(INGESTION / "paste" / "paste.txt")], one_model=True)
    multi = assemble_models(
        [
            read_path(INGESTION / "multi" / "part1.sas"),
            read_path(INGESTION / "multi" / "part2.sas"),
        ],
        one_model=True,
    )
    zipped = assemble_models(unzip_project(INGESTION / "zip" / "model_project.zip"), one_model=True)

    ref = scan_of(single)
    assert scan_of(paste) == ref
    assert scan_of(multi) == ref
    assert scan_of(zipped) == ref
    assert "sales.orders" in ref[0] and "mart.daily_summary" in ref[1]


def test_pdf_ingestion_raises_lossy_issue(_binaries: dict[str, str]) -> None:
    from modeltracex.ingestion import read_path

    artifact = read_path(INGESTION / "pdf" / "model_lossy.pdf")
    assert any("lossy" in i.message.lower() for i in artifact.issues)


def test_exit_sas_python_project_to_docx_and_xlsx(tmp_path: Path) -> None:
    """[EXIT] A SAS+Python project → valid Part A DOCX + Part B XLSX, headless."""
    from modeltracex.cli import run_project

    m1 = json.dumps(
        {
            "purpose": "Build staging.",
            "input_tables": [{"name": "raw.customer_events", "columns": ["cust_id"]}],
            "output_tables": [{"name": "work.staging", "columns": ["cust_id", "amount_net"]}],
            "lineage_rows": ["raw.customer_events.cust_id -> work.staging.cust_id"],
        }
    )
    m2 = json.dumps(
        {
            "purpose": "Score.",
            "input_tables": [{"name": "work.staging"}],
            "output_tables": [{"name": "mart.customer_scores"}],
            "lineage_rows": ["work.staging.amount_net -> mart.customer_scores.total_net"],
        }
    )
    provider = FakeProvider(responses={"m1_build_staging": m1, "m2_score": m2}, default="{}")
    sources = [
        str(CORPUS / "projects" / "stitch_scoring" / "m1_build_staging.sas"),
        str(CORPUS / "projects" / "stitch_scoring" / "m2_score.sas"),
        str(CORPUS / "languages" / "python" / "feature_scoring.py"),  # the Python model
    ]
    state, outputs = run_project(sources, str(tmp_path), provider)

    assert {m.language for m in state.models} >= {Language.SAS, Language.PYTHON}
    docx = [o for o in outputs if o.endswith(".docx")]
    xlsx = [o for o in outputs if o.endswith(".xlsx")]
    assert docx and xlsx

    doc = Document(docx[0])
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 1"] == SECTION_TITLES
    wb = load_workbook(xlsx[0])
    assert wb.sheetnames == [name for name, _ in SHEETS]
